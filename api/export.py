"""生职迹 Career Trail · 简历导出
把库里的简历渲染成可投递的文件：
- GET /api/resumes/{id}/export/pdf  —— reportlab 排版，嵌入 NotoSansSC 中文字体
- GET /api/resumes/{id}/export/docx —— python-docx 生成 Word

**五套版式模板**（与前端编辑器里的「模板」一一对应，见 src/lib/resume-templates.ts）：
    formal   经典求职：黑白一栏式，org左/职位中/日期右三段式条目行，要点冒号前加粗
    classic  经典居中：墨蓝 + 细线分区，最稳妥的正式版面
    band     现代色带：紫罗兰色带页头 + 实色分区标题，信息型选手
    timeline 时间轴：青蓝 + 左侧色条分区标题，时间右对齐，经历多时最清晰
    bold     醒目大字：热粉 + 超大姓名，设计/内容岗的敲门砖

顺序与前端模板选择器一致（formal 排第一）。DEFAULT_TEMPLATE 仍是 classic：
没存过 template 的老简历保持原样，不因排序调整而改版式。

照片（content.photo 的 data URL）会按模板位置嵌进页头并做圆角裁切。
"""
from __future__ import annotations

import base64
import io
import json
import re
import urllib.parse
from pathlib import Path

from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from reportlab.pdfgen import canvas
from reportlab.platypus import BaseDocTemplate, SimpleDocTemplate

from .db import get_db

# ---------------- 中文避头点 ----------------
# reportlab 自带的 CJK 避头点表（ALL_CANNOT_START）只有日文习惯的一小部分
# （、。」等），缺全角逗号、冒号、分号、问号、叹号、右引号、右括号这些
# 中文高频标点 —— 不补的话，自动换行会把「，」「：」折到行首。
# 这张表在 reportlab 里被按值绑定到了两个模块（lib.textsplit 与
# platypus.paragraph），两处都要补。
import reportlab.lib.textsplit as _rl_textsplit  # noqa: E402
import reportlab.platypus.paragraph as _rl_paragraph  # noqa: E402

_KINSOKU_EXTRA = "，：；！？·…—–》〉」』】〉”’％‰%、。+"
_rl_textsplit.ALL_CANNOT_START += _KINSOKU_EXTRA
_rl_paragraph.ALL_CANNOT_START += _KINSOKU_EXTRA

router = APIRouter(prefix="/api/resumes", tags=["export"])

FONT_DIR = Path(__file__).resolve().parent.parent / "assets" / "fonts"

# ---------------- 模板皮肤 ----------------
# accent 主色 / header 页头形态 / section 分区标题形态
TEMPLATES: dict[str, dict[str, str]] = {
    "formal": {"accent": "#111827", "muted": "#4B5563", "header": "center", "section": "rule"},
    "classic": {"accent": "#1E3A5F", "muted": "#475569", "header": "center", "section": "rule"},
    "band": {"accent": "#6D28D9", "muted": "#6B7280", "header": "band", "section": "block"},
    "timeline": {"accent": "#0E7490", "muted": "#4B5563", "header": "left", "section": "bar"},
    "bold": {"accent": "#BE185D", "muted": "#52525B", "header": "big", "section": "big"},
}
DEFAULT_TEMPLATE = "classic"


def skin_of(content: dict) -> dict[str, str]:
    tid = str(content.get("template") or DEFAULT_TEMPLATE)
    tpl = dict(TEMPLATES.get(tid, TEMPLATES[DEFAULT_TEMPLATE]))
    tpl["_id"] = tid  # 个别版式（formal）需要知道自己是谁来做差异化渲染
    return tpl


# ---------------- 输出前的清洗 ----------------

# 条目标题超过这个字数就该独占一行：三段式条目行的首栏只有 0.36 幅宽（约 64mm），
# 10pt 中文一行约 18 字，再长就会折成两三行、把职位与日期挤在旁边。
LONG_HEADING_LIMIT = 16
# 单栏版式（时间轴）标题左、时间右，首栏约 38 字
LONG_HEADING_LIMIT_WIDE = 30

# 每个板块标题前的留白（mm）——「实习经历」这类栏目标题前空一小行
SECTION_GAP_MM = 4


def _item_has_content(it: dict) -> bool:
    return bool(
        str(it.get("heading") or "").strip()
        or str(it.get("subheading") or "").strip()
        or any(str(b).strip() for b in (it.get("bullets") or []))
    )


def output_sections(content: dict) -> list[dict]:
    """导出用的板块列表：**空板块直接丢掉**。

    生成简历时模型常留下一个只有标题的板块（如「科研经历」下面什么都没有），
    打印出来就是一个孤零零的标题；板块内被清空的条目同理。
    """
    out: list[dict] = []
    for sec in content.get("sections") or []:
        if not isinstance(sec, dict):
            continue
        name = str(sec.get("name") or "").strip()
        items = [it for it in (sec.get("items") or []) if isinstance(it, dict) and _item_has_content(it)]
        if name and items:
            out.append({"name": name, "items": items})
    return out


def _load_resume(resume_id: int) -> dict:
    with get_db() as conn:
        row = conn.execute("SELECT * FROM resumes WHERE id = ?", (resume_id,)).fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="简历不存在")
    d = dict(row)
    d["content"] = json.loads(d["content"])
    return d


def _safe_filename(title: str) -> str:
    name = re.sub(r'[\\/:*?"<>|]', "_", title).strip() or "resume"
    return name


def _content_disposition(filename: str, ext: str) -> str:
    quoted = urllib.parse.quote(f"{filename}.{ext}")
    return f"attachment; filename*=UTF-8''{quoted}"


def decode_photo(data_url: str) -> bytes | None:
    """data URL → 图片字节。只接受 image/* 的 base64，异常一律当作没有照片。

    前端统一压成 JPEG（reportlab 能原生读 JPEG，不依赖 Pillow）。
    """
    s = (data_url or "").strip()
    if not s.startswith("data:image/") or "," not in s:
        return None
    try:
        return base64.b64decode(s.split(",", 1)[1], validate=False)
    except Exception:
        return None


# ---------------- PDF ----------------
class _FlushDoc(SimpleDocTemplate):
    """Frame 内边距归零的文档模板。

    SimpleDocTemplate 默认给 Frame 留 6pt 内边距：**纯段落**整体右移 6pt，
    而 Table（宽度=可用宽度、hAlign 居中）恰好回到页边距 —— 两种容器起点差 6pt，
    这就是「栏目标题比下面的条目内容缩进一点」的根源。内边距归零后
    段落、表格、分隔线共用同一条左边界 = 页边距。
    """

    def build(self, flowables, onFirstPage=None, onLaterPages=None, canvasmaker=None):  # noqa: N802
        from reportlab.platypus import Frame, PageTemplate

        self._calc()
        frame = Frame(
            self.leftMargin, self.bottomMargin, self.width, self.height,
            leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0, id="normal",
        )
        # 页面装饰交给 PageTemplate（我们不用 onFirstPage/onLaterPages 钩子）
        self.addPageTemplates([
            PageTemplate(id="First", frames=[frame], pagesize=self.pagesize),
            PageTemplate(id="Later", frames=[frame], pagesize=self.pagesize),
        ])
        BaseDocTemplate.build(self, flowables, canvasmaker=canvasmaker or canvas.Canvas)


def build_pdf(resume: dict) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib.colors import HexColor, white
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        Flowable, HRFlowable, KeepTogether, Paragraph, Spacer, Table, TableStyle,
    )

    pdfmetrics.registerFont(TTFont("NotoSC", str(FONT_DIR / "NotoSansSC-Regular.ttf")))
    pdfmetrics.registerFont(TTFont("NotoSC-Bold", str(FONT_DIR / "NotoSansSC-Bold.ttf")))

    c = resume["content"]
    tpl = skin_of(c)
    accent = HexColor(tpl["accent"])
    muted = HexColor(tpl["muted"])

    class RoundedImage(Flowable):
        """圆角裁切的照片。黏土风/成品简历里都不该出现硬直角。"""

        def __init__(self, data: bytes, width: float, height: float, radius: float = 3 * mm):
            Flowable.__init__(self)
            self.data = data
            self.width = width
            self.height = height
            self.radius = radius

        def draw(self):
            canv = self.canv
            canv.saveState()
            path = canv.beginPath()
            path.roundRect(0, 0, self.width, self.height, self.radius)
            canv.clipPath(path, stroke=0, fill=0)
            canv.drawImage(
                ImageReader(io.BytesIO(self.data)), 0, 0, self.width, self.height,
                preserveAspectRatio=True, anchor="c", mask="auto",
            )
            canv.restoreState()

    photo_bytes = decode_photo(c.get("photo", ""))
    # 一寸照比例（约 5:7）。21x28mm 在 A4 上足够醒目又不抢正文的风头
    photo_w, photo_h = 21 * mm, 28 * mm

    def photo_flowable():
        if not photo_bytes:
            return None
        # formal 是"最像纸质简历"的版式，照片也照纸质惯例用近乎直角的裁切
        radius = 1 * mm if tpl.get("_id") == "formal" else 3 * mm
        return RoundedImage(photo_bytes, photo_w, photo_h, radius)

    from xml.sax.saxutils import escape as _esc

    def bullet_markup(text: str) -> str:
        """要点 → Paragraph 安全文本。formal 版式额外把「前缀：」加粗（如「AI 产品策略：」），
        这正是纸质简历最常见的写法；其余版式维持原样。"""
        raw = str(text)
        if tpl.get("_id") != "formal":
            return _esc(raw)
        m = re.match(r"^(.{1,16}?)[:：]\s*(.+)$", raw, re.S)
        if not m:
            return _esc(raw)
        return f"<b>{_esc(m.group(1))}：</b>{_esc(m.group(2))}"

    # ---- 文字样式（随模板微调字号；整体压小一档，A4 一页能装更多内容） ----
    # wordWrap="CJK"：中文按字符断行并遵守「避头点」——逗号、句号、顿号这类
    # 标点不会出现在行首（reportlab 会把它挪回上一行行尾），行首不再有标点。
    name_size = {"center": 17, "band": 18, "left": 17.5, "big": 21}[tpl["header"]]
    s_name = ParagraphStyle(
        "name", fontName="NotoSC-Bold", fontSize=name_size, leading=name_size * 1.3,
        textColor=white if tpl["header"] == "band" else accent,
        alignment=1 if tpl["header"] == "center" else 0, wordWrap="CJK",
    )
    # 页头第二行：联系方式（邮箱 · 手机号），紧跟姓名之后
    s_contact = ParagraphStyle(
        "contact", fontName="NotoSC", fontSize=9.5, leading=13.5,
        textColor=white if tpl["header"] == "band" else muted,
        alignment=1 if tpl["header"] == "center" else 0, wordWrap="CJK",
    )
    s_name_band = ParagraphStyle("nameband", parent=s_name, textColor=white)
    s_section = ParagraphStyle("section", fontName="NotoSC-Bold", fontSize=11, leading=15)
    s_heading = ParagraphStyle(
        "heading", fontName="NotoSC-Bold", fontSize=10, leading=14, spaceBefore=5, wordWrap="CJK",
    )
    s_heading_big = ParagraphStyle(
        "headingbig", parent=s_heading, fontSize=10.5, textColor=accent,
    )
    s_sub = ParagraphStyle("sub", fontName="NotoSC", fontSize=8.5, leading=12, textColor=muted, wordWrap="CJK")
    s_sub_right = ParagraphStyle("subright", parent=s_sub, alignment=2)
    s_bullet = ParagraphStyle(
        # 要点文字与栏目标题/条目标题**同一条左边界**（leftIndent=0），
        # 圆点悬挂到页边距里（bulletIndent=-9）——整份简历的每行文字左侧对齐。
        "bullet", fontName="NotoSC", fontSize=9.5, leading=14,
        leftIndent=0, bulletIndent=-9, spaceBefore=1, wordWrap="CJK",
        bulletFontName="NotoSC", bulletFontSize=9.5, bulletColor=accent,
    )

    def formal_row(org_s: str, role_s: str, sub_s: str, role_style: ParagraphStyle):
        """经典求职的三段式条目行：机构左 · 职位中（压轴对称线）· 日期右。

        左右两栏等宽，中间那栏的中心正好是幅宽的一半 = 纸面的中轴。
        「机构」传空串时它仍占位，所以长标题折行后职位也还在轴线上。
        """
        row = Table(
            [[Paragraph(org_s, s_heading), Paragraph(role_s, role_style), Paragraph(sub_s, s_sub_right)]],
            colWidths=[usable * 0.36, usable * 0.28, usable * 0.36],
        )
        row.setStyle(TableStyle([
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ("VALIGN", (0, 0), (-1, -1), "BOTTOM"),
        ]))
        return row

    buf = io.BytesIO()
    doc = _FlushDoc(
        buf, pagesize=A4,
        leftMargin=16 * mm, rightMargin=16 * mm,
        topMargin=14 * mm, bottomMargin=14 * mm,
        title=c.get("title", "简历"),
    )
    usable = doc.width

    title = c.get("title", "简历")
    # 页头第一行是姓名、第二行是联系方式（邮箱 · 手机号）；
    # 老简历没填过姓名时回落到 title，不至于开天窗。
    display_name = str(c.get("name") or "").strip() or title
    contact = str(c.get("contact") or "").strip()
    story: list = []

    # ---- 页头（照片按模板位置放） ----
    if tpl["header"] == "band":
        head_cell: list = [
            Paragraph(display_name, s_name_band),
            Spacer(1, 1.5 * mm),
        ]
        if contact:
            head_cell += [Paragraph(contact, s_contact), Spacer(1, 1 * mm)]
        inner = Table([[head_cell, photo_flowable() or ""]], colWidths=[usable - photo_w - 4 * mm, photo_w + 4 * mm])
        inner.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ]))
        band = Table([[inner]], colWidths=[usable])
        band.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), accent),
            ("LEFTPADDING", (0, 0), (-1, -1), 6 * mm),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6 * mm),
            ("TOPPADDING", (0, 0), (-1, -1), 5 * mm),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5 * mm),
        ]))
        story += [band, Spacer(1, 4 * mm)]
    else:
        head_cell = [Paragraph(display_name, s_name)]
        if contact:
            head_cell += [Spacer(1, 1 * mm), Paragraph(contact, s_contact)]
        if photo_bytes:
            gutter = photo_w + 5 * mm
            if tpl["header"] == "center":
                # 居中页头 + 证件照：左右各留一块**等宽**空白，
                # 居中的姓名才落在纸面的轴对称线上（只留右侧会把姓名挤到偏左）。
                cols: list = ["", head_cell, photo_flowable()]
                widths = [gutter, usable - 2 * gutter, gutter]
            else:
                cols = [head_cell, photo_flowable()]
                widths = [usable - gutter, gutter]
            head = Table([cols], colWidths=widths)
            head.setStyle(TableStyle([
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ]))
            story.append(head)
        else:
            story += head_cell
        # 页头分隔线：formal（经典求职）不画 —— 证件照就贴在页头右侧，
        # 通栏黑线会从照片底下穿过去，破坏一栏式简历的周正感。
        # （不留额外留白：每个板块标题前统一有 SECTION_GAP_MM）
        if tpl["_id"] != "formal":
            thickness = 2.4 if tpl["header"] == "big" else 1
            story += [Spacer(1, 2.5 * mm), HRFlowable(width="100%", thickness=thickness, color=accent, spaceAfter=2)]

    # ---- 正文（空板块已经在 output_sections 里丢掉） ----
    for sec in output_sections(c):
        sec_name = sec["name"]
        gap = Spacer(1, SECTION_GAP_MM * mm)
        # 板块标题的 flowables 先攒着，与第一条要点一起 KeepTogether ——
        # 否则分页时标题可能孤零零落在上一页页底
        head_block: list = [gap]
        if tpl["section"] == "block":
            bar = Table([[Paragraph(sec_name, ParagraphStyle(
                "sb", fontName="NotoSC-Bold", fontSize=10.5, leading=14, textColor=white, wordWrap="CJK",
            ))]], colWidths=[usable])
            bar.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), accent),
                ("LEFTPADDING", (0, 0), (-1, -1), 4 * mm),
                ("TOPPADDING", (0, 0), (-1, -1), 1.4 * mm),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 1.4 * mm),
            ]))
            head_block += [bar, Spacer(1, 1.8 * mm)]
        elif tpl["section"] == "bar":
            cell = Table([["", Paragraph(sec_name, ParagraphStyle(
                "sb2", fontName="NotoSC-Bold", fontSize=11, leading=15, textColor=accent, wordWrap="CJK",
            ))]], colWidths=[1.6 * mm, usable - 1.6 * mm])
            cell.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (0, 0), accent),
                ("LEFTPADDING", (0, 0), (0, 0), 0),
                ("RIGHTPADDING", (0, 0), (0, 0), 0),
                ("LEFTPADDING", (1, 0), (1, 0), 3 * mm),
                ("TOPPADDING", (0, 0), (-1, -1), 0),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]))
            head_block += [cell, Spacer(1, 1.8 * mm)]
        elif tpl["section"] == "big":
            head_block += [
                Paragraph(sec_name, ParagraphStyle(
                    "sb3", fontName="NotoSC-Bold", fontSize=12.5, leading=17, textColor=accent, wordWrap="CJK",
                )),
                HRFlowable(width="100%", thickness=1.6, color=accent, spaceAfter=2),
            ]
        else:
            head_block += [
                Paragraph(sec_name, ParagraphStyle(
                    "sb4", fontName="NotoSC-Bold", fontSize=11, leading=15, textColor=accent, wordWrap="CJK",
                )),
                HRFlowable(width="100%", thickness=1, color=accent, spaceAfter=2),
            ]

        item_blocks: list[list] = []
        for item in sec["items"]:
            block: list = []
            heading = item.get("heading", "")
            sub = item.get("subheading", "")
            if tpl["_id"] == "formal":
                # 经典求职：机构居左、职位居中、日期居右，三段式一行（对应纸质简历惯例）。
                # heading 约定用「机构 · 职位」书写，没有「·」时整段都算机构。
                # 左右两栏等宽 → 中间那栏正好压在纸面的轴对称线上。
                parts = [p.strip() for p in str(heading).split("·") if p.strip()]
                if parts or sub:
                    org = parts[0] if parts else ""
                    role = " · ".join(parts[1:]) if len(parts) > 1 else ""
                    heading_c = ParagraphStyle("headingc", parent=s_heading, alignment=1)
                    if len(org) > LONG_HEADING_LIMIT:
                        # 长标题（比赛 / 论文 / 项目的完整名称）独占一行：
                        # 塞进 0.36 幅宽的首栏会折成两三行，职位和日期被挤在旁边。
                        # 职位仍放中栏，所以依旧落在轴对称线上。
                        block.append(Paragraph(org, s_heading))
                        if role or sub:
                            block.append(formal_row("", role, sub, heading_c))
                    else:
                        block.append(formal_row(org, role, sub, heading_c))
            elif tpl["section"] == "bar":
                # 时间轴：标题左、时间右对齐在同一行；标题过长时日期另起一行
                if heading or sub:
                    if len(str(heading)) > LONG_HEADING_LIMIT_WIDE:
                        block.append(Paragraph(heading, s_heading))
                        if sub:
                            block.append(Paragraph(sub, s_sub_right))
                    else:
                        row = Table(
                            [[Paragraph(heading, s_heading), Paragraph(sub, s_sub_right)]],
                            colWidths=[usable - 42 * mm, 42 * mm],
                        )
                        row.setStyle(TableStyle([
                            ("LEFTPADDING", (0, 0), (-1, -1), 0),
                            ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                            ("VALIGN", (0, 0), (-1, -1), "BOTTOM"),
                        ]))
                        block.append(row)
            else:
                if heading:
                    block.append(Paragraph(heading, s_heading_big if tpl["section"] == "big" else s_heading))
                if sub:
                    block.append(Paragraph(sub, s_sub))
            for b in item.get("bullets", []):
                block.append(Paragraph(bullet_markup(b), s_bullet, bulletText="•"))
            if block:
                item_blocks.append(block)

        # 标题 + 第一条同页；其余条目各自成块
        if item_blocks:
            story.append(KeepTogether(head_block + item_blocks[0]))
            for b in item_blocks[1:]:
                story.append(KeepTogether(b))

    doc.build(story)
    return buf.getvalue()


# ---------------- Word ----------------
def build_docx(resume: dict) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    def set_cjk(run, font="NotoSansSC"):
        run.font.name = font
        run._element.rPr.rFonts.set(qn("w:eastAsia"), font)

    def shade(element_pr, fill_hex: str):
        """给段落/单元格加底色（w:shd），Word 里实现"色带"。"""
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill_hex)
        element_pr.append(shd)

    def left_bar(paragraph, fill_hex: str, size_eighth_pt: str = "30"):
        """给段落加左侧竖条（w:pBdr left）—— Word 里的"时间轴色条"。"""
        p_pr = paragraph._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), size_eighth_pt)
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), fill_hex)
        bdr.append(left)
        p_pr.append(bdr)

    # w:pBdr 在 pPr 里的后继元素（ECMA-376 的 _tag_seq 顺序）：底边框必须插在
    # 这些元素之前，否则 spacing/shd 等在前的段落属性会把边框挤到错误位置
    _PPR_AFTER_PBDR = (
        "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
        "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
        "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
        "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
        "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
        "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
    )

    def bottom_rule(paragraph, fill_hex: str, size_eighth_pt: str = "8", space: str = "2"):
        """给段落加通栏底边框（w:pBdr bottom）—— Word 里的"分区细线"。

        对应 PDF 里的 HRFlowable：页头分隔线、栏目标题下的横线都靠它。
        size_eighth_pt 用 1/8 pt 计（"8" = 1pt 线，"19" ≈ 2.4pt 粗线）。
        """
        p_pr = paragraph._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), size_eighth_pt)
        bottom.set(qn("w:space"), space)
        bottom.set(qn("w:color"), fill_hex)
        bdr.append(bottom)
        p_pr.insert_element_before(bdr, *_PPR_AFTER_PBDR)

    def set_col_widths(table, widths_mm: list):
        """按给定宽度（mm）固定表格各列。

        Word 的固定版式实际以 w:tblGrid 为准，只改 cell.width 会被忽略
        （三栏等宽的表就会退化成"三等分"，中间的姓名反而被推到偏右）。
        """
        table.autofit = False
        grid = table._tbl.find(qn("w:tblGrid"))
        if grid is not None:
            for col, w in zip(grid.findall(qn("w:gridCol")), widths_mm):
                col.set(qn("w:w"), str(int(Mm(w).twips)))
        for row in table.rows:
            for cell, w in zip(row.cells, widths_mm):
                cell.width = Mm(w)

    c = resume["content"]
    tpl = skin_of(c)
    accent_rgb = RGBColor.from_string(tpl["accent"].lstrip("#"))
    gray_rgb = RGBColor.from_string(tpl["muted"].lstrip("#"))
    accent_fill = tpl["accent"].lstrip("#").upper()

    doc = Document()
    style = doc.styles["Normal"]
    style.font.size = Pt(10)
    style.font.name = "NotoSansSC"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "NotoSansSC")
    # 段落节奏全部显式控制：python-docx 模板的 docDefaults 自带
    # 「段后 10pt + 1.15 倍行距」，不清零的话每段之间都凭空多出一小行
    # （两段实习经历之间、要点之间全被撑开），与 PDF 的紧凑排版对不上。
    style_pf = style.paragraph_format
    style_pf.space_before = Pt(0)
    style_pf.space_after = Pt(0)
    style_pf.line_spacing = 1.0

    # 页面与 PDF 对齐：A4 + 16/14mm 页边距。
    # python-docx 的默认版式是 Letter，正文宽度与 PDF 差 12mm，
    # 于是同一份简历的 Word 版看起来总是"宽一点、松一点"，排版对不上。
    page = doc.sections[0]
    page.page_width = Mm(210)
    page.page_height = Mm(297)
    page.left_margin = page.right_margin = Mm(16)
    page.top_margin = page.bottom_margin = Mm(14)
    content_w = Mm(178)  # 210 - 16 × 2

    title = c.get("title", "简历")
    # 页头第一行姓名、第二行联系方式（与 PDF 同一结构）；没填姓名回落到 title
    display_name = str(c.get("name") or "").strip() or title
    contact = str(c.get("contact") or "").strip()
    photo = decode_photo(c.get("photo", ""))
    name_size = {"center": 17, "band": 18, "left": 17.5, "big": 21}[tpl["header"]]
    is_formal = tpl.get("_id") == "formal"
    # 条目标题前的呼吸感：与 PDF 的 s_heading.spaceBefore=5 同一节奏
    ITEM_GAP = Pt(5)

    def formal_para(org_s: str, role_s: str, sub_s: str):
        """经典求职的三段式条目行（制表位实现）。

        居中制表位取正文宽度的一半，所以职位永远压在纸面的轴对称线上；
        「机构」传空串时它不占文字位，但两个制表位是绝对位置，职位照样居中。
        """
        hp = doc.add_paragraph()
        hp.paragraph_format.space_before = ITEM_GAP
        hp.paragraph_format.tab_stops.add_tab_stop(Mm(89), WD_TAB_ALIGNMENT.CENTER)
        hp.paragraph_format.tab_stops.add_tab_stop(content_w, WD_TAB_ALIGNMENT.RIGHT)
        r = hp.add_run(org_s)
        r.bold = True
        set_cjk(r)
        if role_s:
            r = hp.add_run("\t" + role_s)
            r.bold = True
            set_cjk(r)
        if sub_s:
            r2 = hp.add_run("\t" + sub_s)
            r2.font.size = Pt(9)
            r2.font.color.rgb = gray_rgb
            set_cjk(r2)
        return hp

    def add_contact(paragraph_parent, align):
        """页头第二行：联系方式（邮箱 · 手机号）。没填就不占行。"""
        if not contact:
            return
        cp = paragraph_parent.add_paragraph()
        cp.alignment = align
        cr = cp.add_run(contact)
        cr.font.size = Pt(9.5)
        cr.font.color.rgb = gray_rgb
        set_cjk(cr)

    # ---- 页头 ----
    if tpl["header"] == "band":
        tbl = doc.add_table(rows=1, cols=2)
        left, right = tbl.rows[0].cells
        shade(left._tc.get_or_add_tcPr(), accent_fill)
        shade(right._tc.get_or_add_tcPr(), accent_fill)
        p = left.paragraphs[0]
        r = p.add_run(display_name)
        r.bold = True
        r.font.size = Pt(name_size)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cjk(r)
        if contact:
            pc = left.add_paragraph()
            rc = pc.add_run(contact)
            rc.font.size = Pt(9.5)
            rc.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            set_cjk(rc)
        if photo:
            right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            right.paragraphs[0].add_run().add_picture(io.BytesIO(photo), width=Mm(20))
        gap_p = doc.add_paragraph()
        gap_p.paragraph_format.space_after = Pt(0)
        gap_p.add_run("").font.size = Pt(4)  # 色带与分隔线之间的小间隔（PDF Spacer 4mm 由下方线段落补足）
    else:
        if photo:
            centered = tpl["header"] == "center"
            tbl = doc.add_table(rows=1, cols=3 if centered else 2)
            cells = tbl.rows[0].cells
            if centered:
                # 左右各留一块等宽空白 → 居中的姓名落在纸面的轴对称线上
                left, right = cells[1], cells[2]
                widths = [22, 134, 22]
            else:
                left, right = cells[0], cells[1]
                widths = [134, 44]
            set_col_widths(tbl, widths)
            hp = left.paragraphs[0]
            hp.alignment = WD_ALIGN_PARAGRAPH.CENTER if tpl["header"] == "center" else WD_ALIGN_PARAGRAPH.LEFT
            r = hp.add_run(display_name)
            r.bold = True
            r.font.size = Pt(name_size)
            r.font.color.rgb = accent_rgb
            set_cjk(r)
            add_contact(left, hp.alignment)
            right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            right.paragraphs[0].add_run().add_picture(io.BytesIO(photo), width=Mm(20))
        else:
            h = doc.add_paragraph()
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER if tpl["header"] == "center" else WD_ALIGN_PARAGRAPH.LEFT
            r = h.add_run(display_name)
            r.bold = True
            r.font.size = Pt(name_size)
            r.font.color.rgb = accent_rgb
            set_cjk(r)
            add_contact(doc, h.alignment)

    # 页头分隔线：与 PDF 同一条通栏线（formal 不画，理由同 PDF ——
    # 证件照贴在页头右侧，通栏线会从照片底下穿过去）。
    if tpl["_id"] != "formal":
        hr = doc.add_paragraph()
        hr_pf = hr.paragraph_format
        hr_pf.space_before = Pt(7)   # ≈ PDF 的 Spacer 2.5mm + spaceAfter
        hr_pf.space_after = Pt(0)
        hr_run = hr.add_run("")
        hr_run.font.size = Pt(2)     # 线段落自身高度压到最小
        # big 页头用粗线（PDF thickness 2.4pt → 19/8 pt），其余 1pt 细线
        bottom_rule(hr, accent_fill, "19" if tpl["header"] == "big" else "8", space="0")

    # ---- 正文（空板块已在 output_sections 里丢掉） ----
    for sec in output_sections(c):
        name = sec["name"]
        p = doc.add_paragraph()
        # 栏目标题前空一小行（与 PDF 的 SECTION_GAP_MM 同一节奏）；
        # 与后面的第一条内容保持同页（keep_with_next），标题不孤悬页底
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(name)
        r.bold = True
        set_cjk(r)
        if tpl["section"] == "block":
            # 色带：实色底 + 白字
            r.font.size = Pt(10.5)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            shade(p._p.get_or_add_pPr(), accent_fill)
        elif tpl["section"] == "bar":
            # 时间轴：左侧竖色条 + 主色文字
            r.font.size = Pt(11)
            r.font.color.rgb = accent_rgb
            left_bar(p, accent_fill)
        elif tpl["section"] == "big":
            r.font.size = Pt(12.5)
            r.font.color.rgb = accent_rgb
            # 标题下的通栏线（PDF thickness 1.6pt → 13/8 pt）
            bottom_rule(p, accent_fill, "13", space="2")
        else:
            r.font.size = Pt(11)
            r.font.color.rgb = accent_rgb
            # rule 版式：标题下一条 1pt 细线 —— 与 PDF 的 HRFlowable 同一条
            bottom_rule(p, accent_fill, "8", space="2")

        for item in sec["items"]:
            heading = item.get("heading", "")
            sub = item.get("subheading", "")
            if is_formal:
                # 经典求职：机构左 · 职位居中（制表位）· 日期右（制表位），三段式一行。
                # 居中制表位取正文宽度的一半 → 「战略实习生」正落在纸面的轴对称线上。
                parts = [p.strip() for p in str(heading).split("·") if p.strip()]
                if parts or sub:
                    org = parts[0] if parts else ""
                    role = " · ".join(parts[1:]) if len(parts) > 1 else ""
                    if len(org) > LONG_HEADING_LIMIT:
                        # 长标题独占一行：塞进左侧会折行并把职位、日期挤在旁边
                        tp = doc.add_paragraph()
                        tp.paragraph_format.space_before = ITEM_GAP
                        r = tp.add_run(org)
                        r.bold = True
                        set_cjk(r)
                        if role or sub:
                            formal_para("", role, sub)
                    else:
                        formal_para(org, role, sub)
            elif heading or sub:
                hp = doc.add_paragraph()
                hp.paragraph_format.space_before = ITEM_GAP
                if tpl["section"] == "bar" and sub and len(str(heading)) <= LONG_HEADING_LIMIT_WIDE:
                    hp.paragraph_format.tab_stops.add_tab_stop(content_w, WD_TAB_ALIGNMENT.RIGHT)
                    r = hp.add_run(heading)
                    r.bold = True
                    set_cjk(r)
                    r2 = hp.add_run("\t" + sub)
                    r2.font.size = Pt(9)
                    r2.font.color.rgb = gray_rgb
                    set_cjk(r2)
                else:
                    if heading:
                        r = hp.add_run(heading)
                        r.bold = True
                        if tpl["section"] == "big":
                            r.font.color.rgb = accent_rgb
                        set_cjk(r)
                    if sub:
                        sp = doc.add_paragraph()
                        r = sp.add_run(sub)
                        r.font.size = Pt(9)
                        r.font.color.rgb = gray_rgb
                        set_cjk(r)
            for b in item.get("bullets", []):
                # 要点：圆点悬挂在页边距里（first_line_indent 为负），
                # 文字从左缩进 0 处起笔 —— 与栏目标题 / 条目标题同一条左边界，
                # 制表符把首行文字精确推回到左缩进位置（Word 悬挂缩进的隐式制表位）。
                bp = doc.add_paragraph()
                pf = bp.paragraph_format
                pf.left_indent = Pt(0)
                pf.first_line_indent = Pt(-9)
                pf.space_before = Pt(1)
                pf.space_after = Pt(0)
                # 显式制表位钉在左边界：圆点后的制表符把首行文字精确推回 0，
                # 与后续折行（左缩进 0）同一条左边界
                pf.tab_stops.add_tab_stop(Pt(0))
                dot = bp.add_run("•\t")
                dot.font.size = Pt(9.5)
                dot.font.color.rgb = accent_rgb
                set_cjk(dot)
                if is_formal:
                    # 经典求职：「前缀：」加粗，与 PDF 一致
                    m = re.match(r"^(.{1,16}?)[:：]\s*(.+)$", str(b), re.S)
                    if m:
                        r = bp.add_run(m.group(1) + "：")
                        r.bold = True
                        set_cjk(r)
                        r = bp.add_run(m.group(2))
                        set_cjk(r)
                        continue
                r = bp.add_run(b)
                set_cjk(r)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------- 路由 ----------------
@router.get("/{resume_id}/export/pdf")
def export_pdf(resume_id: int):
    resume = _load_resume(resume_id)
    data = build_pdf(resume)
    filename = _safe_filename(resume["content"].get("title") or resume["title"])
    return StreamingResponse(
        io.BytesIO(data), media_type="application/pdf",
        headers={"Content-Disposition": _content_disposition(filename, "pdf")},
    )


@router.get("/{resume_id}/export/docx")
def export_docx(resume_id: int):
    resume = _load_resume(resume_id)
    data = build_docx(resume)
    filename = _safe_filename(resume["content"].get("title") or resume["title"])
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": _content_disposition(filename, "docx")},
    )
